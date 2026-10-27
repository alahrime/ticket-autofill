"""Inspection rows -> Word pages from report_template.docx, 20 rows per file ("C&G P3.docx")."""
from __future__ import annotations
import copy, json, os
from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, "report_template.docx")
STATE = os.path.join(HERE, "report.json")  # last customer and page number written
PER_PAGE = 20
SIZE = Pt(14)  # the template's font size
ROW_HEIGHT = Pt(21.6)  # 0.3 in, as on the finished pages; the blank template's 0.4 in rows run 20 rows off the page


def load_state(path=STATE):
    try:
        with open(path, encoding="utf8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return {}


def save_state(state, path=STATE):
    with open(path, "w", encoding="utf8") as f:
        json.dump(state, f)


def write_page(rows, customer, page, path, template=TEMPLATE):
    """One page: '<customer>  PAGE<n>' and 'TOTAL <sum>' on one line above the table, then up to PER_PAGE rows.
    Nothing below the table, so a full page never spills onto a second sheet."""
    doc = Document(template)
    table = doc.tables[0]
    while len(table.rows) - 1 < max(len(rows), PER_PAGE):  # the blank template is one row short
        table._tbl.append(copy.deepcopy(table.rows[-1]._tr))
    for tr in table.rows[1:]:
        tr.height = ROW_HEIGHT
    for tr, row in zip(table.rows[1:], rows):
        for cell, value in zip(tr.cells, (row.date, row.car, row.vin6, row.miles, row.amount)):
            p = cell.paragraphs[0]
            for r in p.runs:
                r._r.getparent().remove(r._r)
            p.add_run(value).font.size = SIZE
    for p in doc.paragraphs[:-1]:  # the template's empty lines under the table; Word needs the last one
        p._p.getparent().remove(p._p)
    doc.paragraphs[-1].paragraph_format.space_after = Pt(0)
    for r in doc.paragraphs[-1].runs:
        r._r.getparent().remove(r._r)
    sec = doc.sections[0]
    head = doc.add_paragraph()
    table._tbl.addprevious(head._p)
    head.paragraph_format.tab_stops.add_tab_stop(sec.page_width - sec.left_margin - sec.right_margin, WD_TAB_ALIGNMENT.RIGHT)
    total = sum(float(r.amount or 0) for r in rows)
    head.add_run(f"{customer}  PAGE{page}\tTOTAL   {total:.2f}").font.size = SIZE
    doc.save(path)


def write_report(rows, customer, first_page, path, template=TEMPLATE):
    """Rows sorted by date, split into pages. path is the first page's file; the rest go beside it.
    Returns the list of files written."""
    rows = sorted(rows, key=lambda r: r.sort_key())
    folder, out = os.path.dirname(path), []
    for i in range(0, len(rows), PER_PAGE):
        page = first_page + i // PER_PAGE
        p = path if i == 0 else os.path.join(folder, f"{customer} P{page}.docx")
        write_page(rows[i:i + PER_PAGE], customer, page, p, template)
        out.append(p)
    return out
