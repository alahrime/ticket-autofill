import json, os, sys, tempfile, threading, urllib.request
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from http.server import ThreadingHTTPServer
from docx import Document
import phone
from inspection import Row

SEEN = []


def fake_reader(photo, slow=True):
    assert slow is False  # the phone never waits on the vision model
    SEEN.append(photo.read())
    return Row(date="08/04/2026", car="24 TOYOTA", vin6="555123", miles="", flags={"miles": "mileage not found"}), None


class NoRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, *a):
        return None


def call(base, path, data=None, cookie=""):
    body = data if isinstance(data, bytes) or data is None else json.dumps(data).encode()
    req = urllib.request.Request(base + path, body, {"Cookie": cookie} if cookie else {})
    with urllib.request.build_opener(NoRedirect).open(req) as r:
        return r.read(), r.headers


def status(base, path, data=None, cookie=""):
    try:
        call(base, path, data, cookie)
        return 200
    except urllib.error.HTTPError as e:
        return e.code


def server():
    phone.Handler.pin, phone.Handler.reader = "482913", staticmethod(fake_reader)
    phone.LOG = os.path.join(tempfile.mkdtemp(), "phone.log")
    phone.Handler.fails.clear()
    srv = ThreadingHTTPServer(("127.0.0.1", 0), phone.Handler)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    return srv, f"http://127.0.0.1:{srv.server_address[1]}"


def test_phone_flow():
    srv, base = server()
    try:
        assert b'name="pin"' in call(base, "/")[0]  # PIN page until logged in
        assert status(base, "/read", b"x") == 401 and status(base, "/word", {"rows": []}) == 401
        assert status(base, "/login", b"pin=111111") == 401
        try:
            call(base, "/login", b"pin=482913")
        except urllib.error.HTTPError as e:  # 303 is not followed
            assert e.code == 303
            cookie = e.headers["Set-Cookie"].split(";")[0]
        assert b"Add receipt photos" in call(base, "/", cookie=cookie)[0]

        assert status(base, "/warm", b"", cookie) == 404
        row = json.loads(call(base, "/read", b"\xff\xd8 fake jpeg", cookie)[0])  # the reply is the row
        assert row == {"date": "08/04/2026", "car": "24 TOYOTA", "vin6": "555123", "miles": "",
                       "flags": {"miles": "mileage not found"}, "error": ""}
        assert SEEN[-1] == b"\xff\xd8 fake jpeg"  # handed over in memory, never written to disk

        rows = [{"date": "08/05/2026", "car": "23 HONDA", "vin6": "023948", "miles": "57333"},
                {"date": "08/04/2026", "car": "24 TOYOTA", "vin6": "555123", "miles": "4500"}]
        data, headers = call(base, "/word", {"customer": "C&G", "page": 7, "rows": rows}, cookie)
        assert data[:2] == b"PK" and "C%26G%20P7.docx" in headers["Content-Disposition"]
        path = os.path.join(tempfile.mkdtemp(), "p.docx")
        with open(path, "wb") as f:
            f.write(data)
        d = Document(path)
        assert [c.text for c in d.tables[0].rows[1].cells] == ["08/04/2026", "24 TOYOTA", "555123", "4500", "20.00"]
        assert d.paragraphs[0].text == "C&G  PAGE7	TOTAL   40.00"

        assert status(base, "/word", {"customer": "C&G", "page": 1, "rows": rows * 11}, cookie) == 400  # >20 rows
        assert status(base, "/rows", cookie=cookie) == 404  # nothing is kept on the PC
        with open(phone.LOG, encoding="utf8") as f:
            log = f.read()
        assert "POST /read 200" in log and "TOYOTA" not in log  # timings only, no customer data
    finally:
        srv.shutdown()


def test_pin_lockout():
    srv, base = server()
    try:
        codes = [status(base, "/login", f"pin={p}".encode()) for p in ["1", "2", "3", "4", "5", "482913"]]
        assert codes == [401] * 5 + [429]  # even the right PIN waits out the lockout
    finally:
        srv.shutdown()


def test_load_pin():
    path = os.path.join(tempfile.mkdtemp(), "pin.txt")
    pin = phone.load_pin(path)
    assert len(pin) == 6 and pin.isdigit() and phone.load_pin(path) == pin

if __name__ == "__main__":
    for n, f in list(globals().items()):
        if n.startswith("test_"):
            f()
    print("ok")
