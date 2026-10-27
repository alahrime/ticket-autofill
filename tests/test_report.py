import os, sys, tempfile
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from docx import Document
from inspection import parse, labelled, vin_year, Row
from docx.shared import Pt
from report import write_report, TEMPLATE

# Made-up receipt; VIN has a valid check digit
RECEIPT = ["VIRGINIA STATE POLICE", "Inspection Date: 3/7/2026",
           "Sticker Number: K0000001 Month Insert: 21400000",
           "License No: N/A VIN: 4T1B11AK9RU555123", "Make: TOYOTA Year Built: 2024",
           "Body Type: CAR 4 DOOR Odometer Reading: 32,500", "Total Charges: $ 20.00"]


def test_labelled_two_per_line():
    f = labelled("\n".join(RECEIPT))
    assert f["vin"] == "4T1B11AK9RU555123" and f["make"] == "TOYOTA" and f["year built"] == "2024"


def test_label_value_reply():
    r = parse(["Label: Inspection Date", "Value: 10/27/2026", "Label: Make", "Value: TOYOTA", "Label: VIN",
               "Value: 4T1B11AK9RU555123", "Label: Year Built", "Value: 2024", "Label: Odometer Reading", "Value: 32500"])
    assert (r.date, r.car, r.vin6, r.miles, r.flags) == ("10/27/2026", "24 TOYOTA", "555123", "32500", {})


def test_parse_receipt():
    r = parse(RECEIPT)
    assert (r.date, r.car, r.vin6, r.miles, r.amount) == ("03/07/2026", "24 TOYOTA", "555123", "32500", "20.00")
    assert r.flags == {}


def test_parse_flags():
    r = parse(["Inspection Date: 10/27/2026", "VIN: 4T1B11AK9RU555124", "Make: HONDA", "Year Built: 2023"])
    assert set(r.flags) == {"vin6", "miles"}
    r = parse(["Make: TOYOTA", "Year Built: 2023", "VIN: 4T1B11AK9RU555123", "Odometer Reading: 5"])
    assert r.flags == {"date": "date not found", "car": "VIN says 2024"}


def test_vin_wrong_length_on_receipt():
    r = parse(["VIN: 4T1B11AK9U555123"])  # receipts have been printed with 16 characters
    assert r.vin6 == "555123" and "16 characters" in r.flags["vin6"]
    assert "misread" in parse(["VIN: 4T1B11AK9RU555124"]).flags["vin6"]


def test_date_slash_read_as_one():
    for raw in ("10127/2026", "10/2712026", "10l27/2026", " 10 / 27 / 2026"):
        assert parse([f"Inspection Date: {raw}"]).date == "10/27/2026", raw
    assert "date" in parse(["Inspection Date: 13/45/2026"]).flags


def test_vin_year():
    assert vin_year("4T1B11AK9RU555123") == "2024"
    assert vin_year("1HGCM82633A004352") == "2003"
    assert vin_year("") == ""


def test_write_report_pages():
    rows = [Row(date=f"08/{d:02d}/2026", car="23 NISSAN", vin6=f"{d:06d}", miles="1000") for d in range(25, 0, -1)]
    path = os.path.join(tempfile.mkdtemp(), "C&G P3.docx")
    files = write_report(rows, "C&G", 3, path)
    assert [os.path.basename(f) for f in files] == ["C&G P3.docx", "C&G P4.docx"]
    d = Document(files[0])
    t = d.tables[0]
    assert len(t.rows) == 21 and t.rows[0].cells[0].text == "DATE"
    assert [c.text for c in t.rows[1].cells] == ["08/01/2026", "23 NISSAN", "000001", "1000", "20.00"]
    assert d.paragraphs[0].text == "C&G  PAGE3	TOTAL   400.00"  # one line, above the table
    assert d.element.body[0] is d.paragraphs[0]._p and d.element.body[1] is t._tbl
    assert len(d.paragraphs) == 2 and d.paragraphs[1].text == ""  # nothing under the table to spill over
    assert all(r.height == Pt(21.6) for r in t.rows[1:])  # 20 rows fit on the page
    d = Document(files[1])
    assert d.tables[0].rows[5].cells[0].text == "08/25/2026" and d.tables[0].rows[6].cells[0].text == ""
    assert d.paragraphs[0].text == "C&G  PAGE4	TOTAL   100.00"
    assert len(Document(TEMPLATE).tables[0].rows) == 20  # template untouched

def test_read_receipt_fast_path():
    """A printed receipt is read by Windows OCR alone; the slow model is never called."""
    if sys.platform != "win32":
        return
    import inspection, ocr
    from PIL import Image, ImageDraw, ImageFont
    img = Image.new("RGB", (1400, 700), "white")
    draw, font = ImageDraw.Draw(img), ImageFont.truetype("arial.ttf", 34)
    for i, line in enumerate(["VEHICLE INSPECTION CERTIFICATION", "Inspection Date: 10/27/2026",
                              "License No: N/A          VIN: 4T1B11AK9RU555123",
                              "Make: TOYOTA             Year Built: 2024",
                              "Body Type: CAR 4 DOOR    Odometer Reading: 32500"]):
        draw.text((40, 40 + i * 110), line, fill="black", font=font)
    path = os.path.join(tempfile.mkdtemp(), "top.png")
    img.save(path)
    slow, ocr.ocr_lines = ocr.ocr_lines, lambda *a: (_ for _ in ()).throw(AssertionError("model called"))
    try:
        row, err = inspection.read_receipt(path)
    finally:
        ocr.ocr_lines = slow
    assert err is None and (row.date, row.car, row.vin6, row.miles, row.flags) == (
        "10/27/2026", "24 TOYOTA", "555123", "32500", {})


def test_fast_read_tries_other_sizes():
    import inspection
    fast = inspection.windows_ocr
    tried = []

    def ocr(img, size):  # only the 2400 px read gets the date
        tried.append(size)
        return RECEIPT if size == 2400 else [l for l in RECEIPT if "Date" not in l]
    inspection.windows_ocr = ocr
    try:
        from PIL import Image
        row = inspection.fast_read(Image.new("RGB", (1932, 2576)))
        assert tried == [1600, 2400] and row.flags == {} and row.date == "03/07/2026"
    finally:
        inspection.windows_ocr = fast


def test_read_receipt_falls_back_to_model():
    import inspection, ocr
    from PIL import Image
    fast, slow = inspection.windows_ocr, ocr.ocr_lines
    sizes = []
    inspection.windows_ocr = lambda img, size: sizes.append(size) or ["Make: TOYOTA"]  # OCR missed most fields
    ocr.ocr_lines = lambda img, prompt: RECEIPT
    try:
        row, err = inspection.read_receipt(Image.new("RGB", (10, 10)))
        assert err is None and row.flags == {} and row.vin6 == "555123"
        ocr.ocr_lines = lambda img, prompt: (_ for _ in ()).throw(RuntimeError("Can't reach Ollama"))
        row, err = inspection.read_receipt(Image.new("RGB", (10, 10)))
        assert err == "Can't reach Ollama" and row.car == "TOYOTA" and "date" in row.flags  # OCR's partial row stays
        assert sizes[:1] == [1600]  # a 10x10 photo: tried once at the usual size, no pointless bigger sizes
        ocr.ocr_lines = lambda img, prompt: (_ for _ in ()).throw(AssertionError("model called"))
        row, err = inspection.read_receipt(Image.new("RGB", (10, 10)), slow=False)
        assert err is None and row.car == "TOYOTA" and "date" in row.flags  # flagged straight away, no model
    finally:
        inspection.windows_ocr, ocr.ocr_lines = fast, slow

if __name__ == "__main__":
    for n, f in list(globals().items()):
        if n.startswith("test_"):
            f()
    print("ok")
